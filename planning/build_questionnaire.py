from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = Path(__file__).resolve().parent.parent / 'output' / 'questionnaire'
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / 'Questionnaire_decisions_Keke_Beauty.docx'

sections = [
    ('Périmètre et lancement', [
        ('01', 'Dans quel pays et quelles villes Keke Beauty sera-t-elle lancée en premier ?', 'Détermine les zones, la devise, les opérateurs SMS et les moyens de paiement.'),
        ('02', 'Quelle plateforme faut-il livrer au lancement : web sur téléphone, PWA installable, Android, iPhone ou plusieurs ?', 'Le cahier des charges dit « mobile et/ou web » ; ce choix fixe la charge de réalisation et de test.'),
        ('03', 'Quelle date cible, quel budget et quelle équipe sont disponibles ?', 'Permet de transformer la trajectoire de 12 sprints en calendrier réaliste.'),
        ('04', 'Qui prendra les décisions produit et validera chaque sprint ? Qui jouera le rôle de Scrum Master ?', 'Ces responsabilités sont nécessaires pour prioriser le backlog et accepter les livraisons.'),
    ]),
    ('Établissements et catalogue', [
        ('05', 'Un gérant peut-il gérer plusieurs établissements ? Un établissement peut-il avoir plusieurs gérants ou comptes partenaires ?', 'Fixe les cardinalités du MCD et les droits d’accès.'),
        ('06', 'Quelles catégories, prestations, devises et zones doivent être proposées dès le lancement ?', 'Permet de préparer le dictionnaire des données et le catalogue initial.'),
        ('07', 'Quels documents KYC faut-il demander, qui les valide et que se passe-t-il en cas de refus ?', 'Détermine le traitement administratif et la protection des justificatifs.'),
        ('08', 'Les établissements peuvent-ils être visibles gratuitement ? Quelles fonctions exigent un abonnement ?', 'Sépare le droit à la fiche, à la réservation et aux autres fonctions avancées.'),
    ]),
    ('Réservations', [
        ('09', 'Une réservation porte-t-elle sur un employé, une ressource comme une cabine, ou sur la capacité globale du salon ?', 'Conditionne le MCD et la prévention des doubles réservations.'),
        ('10', 'La durée dépend-elle de la prestation ? Peut-on réserver plusieurs prestations dans un même rendez-vous ?', 'Détermine le calcul des créneaux et les relations du MLD.'),
        ('11', 'Quand un créneau est-il bloqué : dès la demande ou après confirmation du partenaire ? Combien de temps une demande reste-t-elle en attente ?', 'Définit les états, les événements du MCT et les règles de concurrence.'),
        ('12', 'Qui peut annuler ou reprogrammer, jusqu’à quand, et le client doit-il confirmer une nouvelle proposition ?', 'Détermine les transitions et les notifications.'),
        ('13', 'Que deviennent les rendez-vous déjà confirmés si le salon est suspendu ou si son abonnement expire ?', 'Évite les incohérences entre réservation, modération et paiement.'),
    ]),
    ('Abonnements et paiements', [
        ('14', 'Quelles offres, quels prix et quelle devise faut-il appliquer ? Le paiement mensuel peut-il coexister avec un engagement ferme d’un an ?', 'Précise contrat, échéancier et tarifs historiques.'),
        ('15', 'Le contrat se renouvelle-t-il automatiquement ? Quelle période de grâce et quels droits en cas d’impayé ?', 'Définit le cycle de vie de l’abonnement.'),
        ('16', 'Quels moyens de paiement sont indispensables au premier lancement : Wave, Orange Money, MTN, Moov, Visa et Mastercard ?', 'La couverture dépend du pays et du prestataire choisi.'),
        ('17', 'Quel agrégateur de paiement et quel fournisseur SMS sont déjà contractualisés ou préférés ?', 'Permet de lancer les essais avec comptes de test et de chiffrer les coûts.'),
    ]),
    ('Administration et exploitation', [
        ('18', 'Quels motifs autorisent la suspension d’un client ou d’un établissement, et qui peut la décider ?', 'Détermine les rôles administrateurs et la traçabilité.'),
        ('19', 'Pendant combien de temps conserver les pièces d’identité, les rendez-vous et les paiements ?', 'À valider selon les règles applicables au pays de lancement.'),
        ('20', 'Quels indicateurs seront utilisés pour juger le lancement réussi : salons inscrits, recherches, réservations, abonnements ou encaissements ?', 'Fixe les critères de recette et les tableaux de bord.'),
        ('21', 'Combien de salons pilotes participeront à la recette et qui assurera le support après mise en service ?', 'Prépare les tests réels, la formation et le transfert d’exploitation.'),
    ]),
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)

styles = doc.styles
for key, size in [('Normal', 10), ('Title', 18), ('Heading 1', 12)]:
    st = styles[key]
    st.font.name = 'Aptos'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Heading 1'].paragraph_format.space_before = Pt(8)
styles['Heading 1'].paragraph_format.space_after = Pt(7)

doc.add_paragraph('Questionnaire de cadrage Keke Beauty', style='Title')
p = doc.add_paragraph('À remplir avec le commanditaire avant de valider les règles de gestion, le MCD, le calendrier et les fournisseurs. Les questions couvrent les décisions laissées ouvertes par le cahier des charges actuel.')
p.paragraph_format.space_after = Pt(10)
meta = doc.add_table(rows=2, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = 'Table Grid'
meta.cell(0,0).text = 'Nom du répondant :'
meta.cell(0,1).text = 'Fonction :'
meta.cell(1,0).text = 'Date :'
meta.cell(1,1).text = 'Version des réponses :'
doc.add_paragraph('Vous pouvez répondre directement dans le Word. Dans le PDF, utilisez les espaces prévus ou annotez le document. Si une réponse n’est pas encore connue, indiquez le responsable et la date prévue pour la décision.')

for heading, questions in sections:
    doc.add_heading(heading, level=1)
    for num, question, reason in questions:
        p=doc.add_paragraph()
        p.paragraph_format.keep_with_next=True
        p.add_run(num + '  ' + question).bold=True
        r=doc.add_paragraph('Pourquoi cette réponse compte : ' + reason)
        r.paragraph_format.keep_with_next=True
        r.runs[0].font.size=Pt(9)
        ans=doc.add_paragraph('Réponse : __________________________________________________________________________________')
        ans.paragraph_format.space_after=Pt(8)
        ans.paragraph_format.keep_together=True

doc.add_heading('Validation des décisions', level=1)
doc.add_paragraph('Décisions encore ouvertes et responsable : ________________________________________________________')
doc.add_paragraph('Date de validation prévue : _____________________________________________________________________')
doc.add_paragraph('Nom et accord du commanditaire : _________________________________________________________________')
doc.save(DOCX)
print(DOCX)
